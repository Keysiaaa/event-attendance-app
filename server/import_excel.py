import openpyxl
import json
import datetime

excel_file = r'C:\Users\Pongo\Downloads\Data Peserta.xlsx'
out_path = r'C:\Users\Pongo\.gemini\antigravity\scratch\event-attendance-app\server\data\participants.json'

wb = openpyxl.load_workbook(excel_file)
sheet = wb.active

participants = []
current_num = 1001

for row_idx in range(2, sheet.max_row + 1):
    email = sheet.cell(row_idx, 1).value
    name = sheet.cell(row_idx, 2).value
    jabatan = sheet.cell(row_idx, 3).value
    daerah_kerja = sheet.cell(row_idx, 4).value
    phone = sheet.cell(row_idx, 5).value

    if not name and not email:
        continue

    name_clean = str(name).strip() if name else ''
    if not name_clean:
        continue

    barcode = f'EVT-{current_num}'
    current_num += 1

    institution_parts = []
    if jabatan:
        institution_parts.append(str(jabatan).strip())
    if daerah_kerja:
        institution_parts.append(str(daerah_kerja).strip())
    institution = ' - '.join(institution_parts) if institution_parts else '-'

    # Normalize phone format if needed
    phone_clean = str(phone).strip() if phone else '-'
    if phone_clean.endswith('.0'):
        phone_clean = phone_clean[:-2]

    participant = {
        'id': barcode,
        'barcode': barcode,
        'name': name_clean,
        'institution': institution,
        'email': str(email).strip() if email else '-',
        'phone': phone_clean,
        'jabatan': str(jabatan).strip() if jabatan else '-',
        'daerahKerja': str(daerah_kerja).strip() if daerah_kerja else '-',
        'attendance': {
            'registrasi_awal': None,
            'sesi_1': None,
            'sesi_2': None,
            'sesi_3': None,
            'sertifikat': None
        },
        'createdAt': datetime.datetime.now(datetime.timezone.utc).isoformat()
    }
    participants.append(participant)

with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(participants, f, ensure_ascii=False, indent=2)

print(f'Imported {len(participants)} participants successfully!')
